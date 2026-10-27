"""Crawl the public DSU website (dsu.edu.in) and extract text from HTML pages,
documents (PDF, Word, Excel, PowerPoint), and images/screenshots on the fly.

Run from the project folder with the venv active:

    python crawl_website.py                       # crawl and extract pages & files
    python crawl_website.py --max-pages 1500      # crawl more pages

Output:
    data/website/pages.jsonl      one JSON line per page/file: url, title, text

Then run:  python ingest.py
"""
import argparse
import io
import json
import os
import re
import time
from collections import Counter, deque
from pathlib import Path
from urllib import robotparser
from urllib.parse import parse_qsl, urlencode, urljoin, urlparse, urlunparse

import requests
from bs4 import BeautifulSoup
from PIL import Image
import fitz  # PyMuPDF

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    import pytesseract
    pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
except ImportError:
    pytesseract = None

START_URL = "https://www.dsu.edu.in/"
USER_AGENT = "Mozilla/5.0 (compatible; DSU-Student-Assistant-Crawler/1.0; student project)"
OUT_DIR = "data/website"

# Skip only non-document media files, archives, and scripts
SKIP_EXT = (
    ".mp4", ".mp3", ".avi", ".mov", ".zip", ".rar", ".7z",
    ".css", ".js", ".xml", ".json", ".ico", ".swf"
)
SKIP_PARTS = (
    "format=feed", "print=1", "tmpl=component", "/component/mailto", "/login",
    "/register", "logout", "cdn-cgi", "/wp-admin", "javascript:",
)

DOCUMENT_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".pptx", ".jpg", ".jpeg", ".png", ".webp")


# ---------------------------------------------------------------- helpers
def normalize(url, allowed):
    """Clean a URL; return None if it is outside allowed domains or not worth crawling."""
    if any(part in url.lower() for part in SKIP_PARTS):
        return None
    p = urlparse(url)
    if p.scheme not in ("http", "https"):
        return None
    host = p.netloc.lower()
    if host == "dsu.edu.in":
        host = "www.dsu.edu.in"
    if not any(host == d or host.endswith("." + d) for d in allowed):
        return None
    path = p.path or "/"
    if len(path) > 1:
        path = path.rstrip("/")
    if path.lower().endswith(SKIP_EXT):
        return None
    query = [(k, v) for k, v in parse_qsl(p.query) if not k.lower().startswith("utm_")]
    return urlunparse(("https", host, path, "", urlencode(query), ""))


def decode_cfemail(hexstr):
    """Cloudflare hides e-mail addresses as hex; decode them back."""
    try:
        key = int(hexstr[:2], 16)
        return "".join(chr(int(hexstr[i:i + 2], 16) ^ key) for i in range(2, len(hexstr), 2))
    except (ValueError, IndexError):
        return ""


def table_to_text(table, soup):
    rows = []
    for tr in table.find_all("tr"):
        cells = [re.sub(r"\s+", " ", c.get_text(" ", strip=True)) for c in tr.find_all(["th", "td"])]
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    lines = []
    if body and all(len(r) == len(header) for r in body) and len(header) > 1:
        for r in body:
            lines.append(" | ".join(f"{h}: {v}" if h else v for h, v in zip(header, r) if v))
    else:
        lines = [" | ".join(c for c in r if c) for r in rows]
    return "\n".join(lines)


def extract_page(soup, url):
    """Return (title, list_of_text_lines) for one HTML page."""
    title = re.sub(r"\s+", " ", soup.title.get_text(strip=True)) if soup.title else url

    if soup.head:
        soup.head.decompose()

    for el in soup.select("[data-cfemail]"):
        el.replace_with(decode_cfemail(el.get("data-cfemail", "")))
    for a in soup.select('a[href*="email-protection#"]'):
        a.replace_with(decode_cfemail(a["href"].split("#", 1)[1]))

    for tag in soup(["script", "style", "noscript", "svg", "iframe", "select", "option", "form"]):
        tag.decompose()

    for table in reversed(soup.find_all("table")):
        table.replace_with("\n" + table_to_text(table, soup) + "\n")

    lines = []
    for raw in soup.get_text("\n").split("\n"):
        line = re.sub(r"\s+", " ", raw).strip()
        if len(line) >= 2 and line not in ("Prev", "Next", "Search"):
            lines.append(line)
    return title, lines


def extract_file_content(content_bytes, url):
    """Detect file type by extension and extract text dynamically with robust error handling."""
    ext = Path(urlparse(url).path).suffix.lower()
    text_lines = []
    
    try:
        if not content_bytes:
            return ""
            
        if ext == ".pdf":
            with fitz.open(stream=content_bytes, filetype="pdf") as doc:
                for page in doc:
                    page_text = page.get_text()
                    if not page_text.strip() and pytesseract:
                        try:
                            pix = page.get_pixmap()
                            img = Image.open(io.BytesIO(pix.tobytes("png")))
                            page_text = pytesseract.image_to_string(img)
                        except Exception:
                            pass
                    text_lines.append(page_text)
                    
        elif ext in (".jpg", ".jpeg", ".png", ".webp"):
            if pytesseract:
                try:
                    img = Image.open(io.BytesIO(content_bytes))
                    ocr_text = pytesseract.image_to_string(img)
                    text_lines.append(ocr_text)
                except Exception as img_err:
                    print(f"  ! Skipped invalid/corrupted image file {url}: {img_err}")
            else:
                text_lines.append("[Image file - OCR skipped, pytesseract not installed]")
            
        elif ext == ".docx" and DocxDocument:
            doc = DocxDocument(io.BytesIO(content_bytes))
            for para in doc.paragraphs:
                if para.text.strip():
                    text_lines.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_lines.append(row_text)
                        
        elif ext == ".xlsx" and load_workbook:
            wb = load_workbook(io.BytesIO(content_bytes), data_only=True)
            for sheet in wb.worksheets:
                text_lines.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_str = " | ".join(str(cell) for cell in row if cell is not None)
                    if row_str.strip():
                        text_lines.append(row_str)
                        
        elif ext == ".pptx" and Presentation:
            prs = Presentation(io.BytesIO(content_bytes))
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            if para.text.strip():
                                text_lines.append(para.text)
    except Exception as e:
        print(f"  ! Failed parsing file {url}: {e}")
        
    return "\n".join(text_lines)


def strip_boilerplate(pages):
    """Menus/footers repeat on every page. Drop lines that appear on >25% of pages."""
    n = len(pages)
    if n < 12:
        return [{"url": p["url"], "title": p["title"], "text": "\n".join(p["lines"])} for p in pages]
    counts = Counter()
    for p in pages:
        counts.update(set(p["lines"]))
    limit = max(5, int(0.25 * n))

    out = []
    common, seen = [], set()
    for line in pages[0]["lines"]:
        if counts[line] > limit and line not in seen:
            seen.add(line)
            common.append(line)
    if common:
        out.append({
            "url": START_URL,
            "title": "DSU website - common header and footer (contacts, helplines, addresses, quick links)",
            "text": "\n".join(common[:400]),
        })
    for p in pages:
        if p.get("is_file"):
            out.append({"url": p["url"], "title": p["title"], "text": "\n".join(p["lines"])})
            continue
        kept = [l for l in p["lines"] if counts[l] <= limit]
        text = "\n".join(kept)
        if len(text) >= 200:
            out.append({"url": p["url"], "title": p["title"], "text": text})
    return out


def load_robots(session):
    rp = robotparser.RobotFileParser()
    try:
        r = session.get(urljoin(START_URL, "/robots.txt"), timeout=15)
        if r.status_code == 200:
            rp.parse(r.text.splitlines())
            return rp
    except requests.RequestException:
        pass
    rp.parse([])
    return rp


# ---------------------------------------------------------------- crawl
def crawl(args):
    allowed = ["dsu.edu.in"] + args.extra_domains
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    rp = load_robots(session)

    queue, seen = deque([START_URL]), {START_URL}
    pages = []

    try:
        while queue and len(pages) < args.max_pages:
            url = queue.popleft()
            if not rp.can_fetch(USER_AGENT, url):
                continue
            
            path_lower = urlparse(url).path.lower()
            is_file_link = path_lower.endswith(DOCUMENT_EXTENSIONS)

            try:
                r = session.get(url, timeout=25)
            except requests.RequestException as e:
                print(f"  ! failed: {url} ({e.__class__.__name__})")
                continue
            if r.status_code != 200:
                print(f"  ! HTTP {r.status_code}: {url}")
                continue

            if is_file_link:
                extracted_text = extract_file_content(r.content, url)
                lines = [l for l in extracted_text.split("\n") if l.strip()]
                if lines:
                    pages.append({
                        "url": url,
                        "title": f"Document/Image: {os.path.basename(path_lower)}",
                        "lines": lines,
                        "is_file": True
                    })
                    print(f"[{len(pages)}] (File/Image parsed) {url}  ({len(lines)} lines)")
                continue

            ctype = r.headers.get("Content-Type", "").lower()
            if "html" not in ctype:
                continue

            soup = BeautifulSoup(r.text, "html.parser")

            # Collect links from <a> tags
            for a in soup.find_all("a", href=True):
                link = normalize(urljoin(url, a["href"]), allowed)
                if link and link not in seen:
                    seen.add(link)
                    queue.append(link)

            # Collect embedded document/PDF links from <iframe> or <embed> tags
            for frame in soup.find_all(["iframe", "embed"], src=True):
                link = normalize(urljoin(url, frame["src"]), allowed)
                if link and link not in seen:
                    seen.add(link)
                    queue.append(link)

            title, lines = extract_page(soup, url)
            pages.append({"url": url, "title": title, "lines": lines, "is_file": False})
            print(f"[{len(pages)}] {url}  ({len(lines)} lines, queue {len(queue)})")
            time.sleep(args.delay)
    except KeyboardInterrupt:
        print("\nStopped by user - saving what was crawled so far...")

    pages_out = strip_boilerplate(pages)
    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, "pages.jsonl")
    with open(path, "w", encoding="utf-8") as f:
        for p in pages_out:
            f.write(json.dumps(p, ensure_ascii=False) + "\n")
    print(f"\nSaved {len(pages_out)} items to {path}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Crawl dsu.edu.in for the DSU chatbot")
    ap.add_argument("--max-pages", type=int, default=600)
    ap.add_argument("--delay", type=float, default=0.7, help="seconds between requests")
    ap.add_argument("--extra-domains", nargs="*", default=[],
                    help="more domains to crawl, e.g. cdsimer.edu.in dsuonline.com")
    crawl(ap.parse_args())